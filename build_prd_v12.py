from pathlib import Path
import os

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT = ROOT / os.environ.get("PRD_OUTPUT_NAME", "教育局学校问卷调研应用PRD_V1.6_终版.docx")
GREEN = "087A5B"
NAVY = "223847"
MUTED = "6D7D88"
LINE = "D9E2E7"
LIGHT = "EFF8F5"
WHITE = "FFFFFF"


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.72)
section.left_margin = Inches(0.72)
section.right_margin = Inches(0.72)
section.header_distance = Inches(0.32)
section.footer_distance = Inches(0.32)


def set_font(run, size=10.2, bold=False, color=NAVY, name="Microsoft YaHei"):
    run.font.name = name
    fonts = run._element.get_or_add_rPr().rFonts
    fonts.set(qn("w:eastAsia"), name)
    fonts.set(qn("w:ascii"), "Arial")
    fonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


for style_name, size, bold, before, after in [
    ("Normal", 10.2, False, 0, 5),
    ("Title", 26, True, 0, 6),
    ("Heading 1", 16, True, 14, 7),
    ("Heading 2", 12.2, True, 10, 5),
    ("Heading 3", 10.8, True, 7, 4),
    ("List Bullet", 10.2, False, 0, 4),
    ("List Number", 10.2, False, 0, 4),
]:
    style = doc.styles[style_name]
    style.font.name = "Microsoft YaHei"
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(NAVY)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.line_spacing = 1.16


def add_page_number(paragraph):
    paragraph.add_run("第 ")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run = paragraph.add_run()
    run._r.append(begin)
    run._r.append(instruction)
    run._r.append(end)
    paragraph.add_run(" 页")
    for item in paragraph.runs:
        set_font(item, 8, False, MUTED)


header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
set_font(header.add_run("教育局学校问卷调研应用 PRD V1.6"), 8.2, False, MUTED)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_page_number(footer)


def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    node = props.find(qn("w:shd")) or OxmlElement("w:shd")
    node.set(qn("w:fill"), fill)
    if node.getparent() is None:
        props.append(node)


def borders(cell):
    props = cell._tc.get_or_add_tcPr()
    root = props.first_child_found_in("w:tcBorders") or OxmlElement("w:tcBorders")
    if root.getparent() is None:
        props.append(root)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = root.find(qn(f"w:{edge}")) or OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), LINE)
        if node.getparent() is None:
            root.append(node)


def margins(cell, top=80, start=95, bottom=80, end=95):
    props = cell._tc.get_or_add_tcPr()
    root = props.first_child_found_in("w:tcMar") or OxmlElement("w:tcMar")
    if root.getparent() is None:
        props.append(root)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = root.find(qn(f"w:{name}")) or OxmlElement(f"w:{name}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        if node.getparent() is None:
            root.append(node)


def row_no_split(row, header=False):
    props = row._tr.get_or_add_trPr()
    no_split = OxmlElement("w:cantSplit")
    props.append(no_split)
    if header:
        repeat = OxmlElement("w:tblHeader")
        repeat.set(qn("w:val"), "true")
        props.append(repeat)


def cell_text(cell, value, bold=False, color=NAVY, size=8.2, center=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.08
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    set_font(p.add_run(str(value)), size, bold, color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    margins(cell)
    borders(cell)


def table(headers, rows, widths=None, font_size=8.2, center_columns=None):
    result = doc.add_table(rows=1, cols=len(headers))
    result.alignment = WD_TABLE_ALIGNMENT.CENTER
    result.autofit = False
    center_columns = set(center_columns or [])
    row_no_split(result.rows[0], True)
    for index, header_value in enumerate(headers):
        cell_text(result.rows[0].cells[index], header_value, True, WHITE, font_size, True)
        shade(result.rows[0].cells[index], GREEN)
    for row_index, values in enumerate(rows):
        row = result.add_row()
        row_no_split(row)
        for index, value in enumerate(values):
            cell_text(row.cells[index], value, False, NAVY, font_size, index in center_columns)
            if row_index % 2:
                shade(row.cells[index], "F7FAFB")
    if widths:
        for row in result.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)
    return result


def h1(text):
    doc.add_paragraph(text, style="Heading 1")


def h2(text):
    doc.add_paragraph(text, style="Heading 2")


def para(text):
    p = doc.add_paragraph()
    set_font(p.add_run(text))
    return p


def bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    set_font(p.add_run(text))
    return p


def number(text):
    p = doc.add_paragraph(style="List Number")
    set_font(p.add_run(text))
    return p


def note(label, text):
    result = doc.add_table(rows=1, cols=1)
    result.autofit = False
    cell = result.cell(0, 0)
    shade(cell, LIGHT)
    borders(cell)
    margins(cell, 125, 145, 125, 145)
    cell.text = ""
    p = cell.paragraphs[0]
    set_font(p.add_run(f"{label}："), 9.8, True, GREEN)
    set_font(p.add_run(text), 9.8)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def screenshot(path_name, caption):
    path = ROOT / path_name
    if not path.exists():
        return False
    doc.add_picture(str(path), width=Inches(6.65))
    p = doc.paragraphs[-1]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_p = doc.add_paragraph()
    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(caption_p.add_run(caption), 8.6, False, MUTED)
    return True


# 封面
title_kicker = doc.add_paragraph()
title_kicker.paragraph_format.space_before = Pt(22)
title_kicker.paragraph_format.space_after = Pt(5)
set_font(title_kicker.add_run("产品需求文档（PRD）"), 11, True, GREEN)
title = doc.add_paragraph(style="Title")
title.paragraph_format.space_after = Pt(5)
set_font(title.add_run("教育局学校问卷调研应用"), 26, True, "000000")
subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_after = Pt(18)
set_font(subtitle.add_run("部门 / 工作组接收范围回显与发布快照"), 14, False, "000000")
table(["文档属性", "内容"], [
    ["版本", "V1.6"],
    ["状态", "终版 / 可进入研发"],
    ["日期", "2026-09-05"],
    ["适用端", "教育局 PC、学校 PC、教育局移动端、学校移动端"],
    ["承接版本", "教育局学校问卷调研应用 PRD V1.5"],
], [1.35, 5.35], 9.0)
note("产品结论", "完整选择部门或工作组时，汇总页只回显范围名称；部分选择时回显具体人员。编辑态保留动态范围及手动拆分意图，发布确认时按最新有效成员解析、按 personId 去重并固化人员快照。PC 与移动端复用同一数据模型和计算规则。")


h1("1. 背景与目标")
h2("1.1 背景")
para("现有接收对象汇总页把完整部门或工作组展开为人员姓名，创建人无法判断问卷是面向整个组织范围还是一组固定人员。组织成员变化、范围交叉覆盖和再次编辑时，显示结果与原始选择意图存在偏差。")
para("本版本把完整范围与明确人员建模为不同层级：完整范围保持动态范围对象，部分选择保持人员对象；发布时统一转为不可变人员快照，保证已发布问卷的接收对象、应提交人数与统计结果不随组织变化。")
h2("1.2 目标")
bullet("教育局端、学校端的 PC 与移动端使用一致的完整范围识别、手动拆分、交叉去重和回显规则。")
bullet("完整部门或完整工作组仅回显名称；手动选择具体人员或拆分后的部分范围仅回显人员姓名。")
bullet("模块“已选择 N 人”按 personId 去重后的实际人数计算，不以标签数量或范围人数直接相加。")
bullet("草稿与模板保存选择意图；发布确认时使用最新有效成员并固化接收人员快照。")
bullet("为后端明确权限、事务、数据结构、去重、异常与缓存边界。")
h2("1.3 成功指标【假设】")
table(["指标", "目标", "核验方式"], [
    ["四端规则一致率", "100%", "四种端组合逐场景比对"],
    ["模块人数准确率", "100%", "对比实际有效人员 personId 去重集合"],
    ["选择意图恢复率", "100%", "草稿、模板、PC/移动跨端再次编辑"],
    ["发布快照稳定率", "100%", "发布前后修改组织成员并复核快照"],
], [1.55, 1.2, 3.95], 8.5, {1})


h1("2. 用户与使用场景")
h2("2.1 用户")
table(["用户", "端", "任务", "权限边界"], [
    ["教育局问卷创建人", "PC/移动", "选择局内部门、工作组或具体人员", "仅限当前教育局授权组织树"],
    ["学校问卷创建人", "PC/移动", "选择校内部门或具体人员", "仅限当前学校"],
    ["问卷接收人", "PC/移动", "接收并填写已发布问卷", "仅可访问发布快照内本人任务"],
    ["研发与测试", "评审原型", "查看 prd 规则、接口和异常口径", "仅原型评审入口"],
], [1.45, 1.0, 2.75, 2.0], 8.2)
h2("2.2 场景")
bullet("创建人勾选完整“办公室”，汇总页展示“办公室”，右上角人数仍显示该部门有效人员的去重人数。")
bullet("创建人逐人勾选并最终覆盖“项目办”全部有效成员，系统自动把人员集合归一化为“项目办”。")
bullet("创建人从完整范围取消一人，当前范围拆分为剩余人员；其他仍覆盖该人员的完整范围不拆分。")
bullet("创建人重新补齐被拆分范围的全部有效成员，系统清除手动拆分标识并恢复范围名称。")
bullet("同一人员属于多个已选范围时，模块人数只计一次；选择器标识“已由其他范围选中”。")


h1("3. 需求范围")
h2("3.1 In Scope")
bullet("内部部门与工作组的完整范围选择、逐人选择、逐人选满自动归一化、手动拆分与重新补齐。")
bullet("范围交叉覆盖、来源级删除、人员级删除、模块人数去重、PC 与移动端汇总回显。")
bullet("选择器临时副本、确认提交、取消回滚、草稿和模板的原始选择与拆分标识持久化。")
bullet("发布前范围有效性与权限校验、最新有效成员解析、全局 personId 去重和发布快照固化。")
bullet("原型内 prd 规则、事件埋点、自动化测试、交互截图和研发实现约束。")
h2("3.2 Out of Scope")
bullet("学校、家长和班级模块的既有选择规则调整。")
bullet("组织架构、部门、工作组及人员状态的维护流程。")
bullet("已发布问卷接收快照随组织变化自动扩缩。")
bullet("生产环境数据库表名、消息发送、催办、撤回和重新发布。")


h1("4. 功能需求列表")
requirements = [
    ["范围选择", "完整部门", "选择部门当前全部 status=active 且 canReceive=true 的成员时，保存 department scope 对象并仅回显部门名称。", "P0", "人员数按 personId 去重"],
    ["范围选择", "完整工作组", "选择工作组当前全部有效可接收成员时，保存 workgroup scope 对象并仅回显工作组名称。", "P0", "空工作组不可整选"],
    ["范围选择", "部分人员", "未覆盖范围全部有效成员时，保存人员对象并在汇总页回显具体人员姓名。", "P0", "不显示范围名称"],
    ["自动归一化", "逐人选满", "逐人操作后覆盖某范围全部有效成员时，自动移除该范围来源的人员对象并写入 scope 对象。", "P0", "PC与移动一致"],
    ["手动拆分", "取消一人", "从完整范围取消一人时，仅拆分当前范围为剩余人员并记录手动拆分标识；其他完整范围保持不变。", "P0", "不可隐式拆分其他来源"],
    ["手动拆分", "重新补齐", "只有以当前范围来源重新补齐全部有效成员后，才清除手动拆分标识并恢复范围名称。", "P0", "其他范围覆盖不等于补齐"],
    ["交叉覆盖", "多个完整范围", "同一批人员完整覆盖多个部门或工作组时，展示所有完整范围名称；实际人数全局去重。", "P0", "标签数不等于人数"],
    ["交叉覆盖", "其他来源提示", "人员仍被其他完整范围覆盖时，在当前选择器显示“已由其他范围选中”，当前范围勾选态独立计算。", "P0", "避免全局勾选歧义"],
    ["删除", "删除范围标签", "删除范围标签只移除该 scope 来源，不得移除其他范围或明确人员来源。", "P0", "重新计算回显和人数"],
    ["删除", "删除人员标签", "删除人员标签只移除该明确人员对象；由完整范围覆盖的同一人员继续保留。", "P0", "按来源处理"],
    ["计数", "模块人数", "模块“已选择 N 人”取当前范围解析人员与明确人员的 personId 去重集合大小。", "P0", "停用人员不计"],
    ["选择会话", "确认与取消", "打开选择器时创建临时副本；确认替换当前模块选择，取消、关闭、遮罩、Escape 或移动返回均回滚。", "P0", "移动底部按钮为“确认”"],
    ["草稿模板", "语义保存", "保存范围对象、人员对象和手动拆分范围标识；再次编辑不得错误恢复完整范围。", "P0", "需深拷贝"],
    ["范围校验", "有效与授权", "范围不存在、越权、无有效成员或不可选择时标记失效并阻止发布，不得静默转为空。", "P0", "返回可识别错误"],
    ["发布", "即时解析", "发布确认事务内按最新组织目录解析全部 scope，叠加明确人员并按 personId 全局去重。", "P0", "同一权限快照"],
    ["发布", "人员快照", "发布成功后只保存固化人员快照；组织后续变化不得修改该问卷接收对象。", "P0", "记录 snapshotAt"],
    ["权限", "服务端二次鉴权", "服务端根据登录用户、portalType、orgId 和 permissionScope 校验范围与人员，不信任前端提交名称和人数。", "P0", "越权返回403"],
    ["PRD入口", "规则说明", "PC 与移动端接收对象页面右上角 prd 打开只读说明，覆盖模型、算法、接口、异常与验收。", "P1", "生产环境可隐藏"],
]
table(["功能模块", "功能点", "需求描述", "优先级（P0 / P1 / P2）", "备注说明"], requirements, [1.0, 1.05, 3.35, 0.78, 0.67], 7.35, {3})


h1("5. 核心流程与交互说明")
h2("5.1 选择与归一化")
number("进入接收对象步骤，选择内部部门或工作组。系统只加载当前组织授权范围内 status=active 且 canReceive=true 的人员。")
number("勾选完整范围时写入 scope 对象；逐人选择时写入 person 对象，并记录该选择来自哪个 scope。")
number("每次选择变更后按范围成员集合判断完整覆盖。非手动拆分范围若全部成员已被覆盖，则自动归一化为范围对象。")
number("同一人员被多个来源覆盖时保留全部来源关系，最终人数按 personId 去重。")
h2("5.2 手动拆分与恢复")
number("用户在完整范围内取消人员时，只删除当前 scope 对象，将当前范围其余成员写为带当前 scopeId 来源的人员对象，并写入 manualPartialScopeKey。")
number("当前取消人员若仍由其他范围覆盖，保留该覆盖并在当前范围显示“已由其他范围选中”；模块人数按最终并集计算。")
number("用户以当前范围来源补齐全部成员后，删除 manualPartialScopeKey，移除该范围来源的人员记录并恢复 scope 标签。")
h2("5.3 删除与回显")
number("汇总页按模块展示：scope 标签只显示 scopeName；person 标签只显示 name；不在 scope 标签展开人员。")
number("删除 scope 标签只移除该来源；删除人员标签只移除该明确人员。变更后重新归一化并计算模块人数。")
number("PC 与移动端共用同一正式选择 Map；移动端底部“确认”提交临时副本，返回主页面后回显一致。")
h2("5.4 草稿、模板与发布")
number("草稿和模板保存 recipientSelections 与 recipientManualPartialScopes，不提前生成发布人员快照。")
number("再次编辑时恢复范围、明确人员、来源和拆分标识；按当前组织目录标记已失效范围，但不得自动删除。")
number("点击发布后，服务端在同一事务与权限快照内重新校验并解析全部范围，合并明确人员，按 personId 去重。")
number("发布成功保存 recipientSnapshot，包含 personId、姓名快照、来源范围、snapshotAt 与 frozen=true；原始动态范围不再驱动该问卷后续接收对象变化。")
h2("5.5 四端交互参考")
captured = False
captured |= screenshot("recipient-scope-bureau-pc.png", "图 1  教育局端 PC：完整部门与工作组回显范围名称")
captured |= screenshot("recipient-scope-school-pc.png", "图 2  学校端 PC：完整部门回显范围名称")
captured |= screenshot("recipient-scope-bureau-mobile.png", "图 3  教育局移动端：与 PC 共用回显和计数规则")
captured |= screenshot("recipient-scope-school-mobile.png", "图 4  学校移动端：确认后返回接收对象汇总")
captured |= screenshot("recipient-scope-prd.png", "图 5  原型 prd 入口：数据模型与后端规则")
if not captured:
    note("截图占位", "构建时未发现四端交互截图。交付前必须生成并重新构建文档。")


h1("6. 异常场景与边界条件")
table(["异常/边界", "处理规则", "前端反馈"], [
    ["范围无有效成员", "不可选择；已保存范围再次编辑时标记 invalid。", "暂无可接收人员"],
    ["范围被删除或越权", "保留失效标识并阻止发布；不得解析为空数组后继续。", "接收范围已失效，请重新选择"],
    ["人员停用或不可接收", "不参与完整范围判断、人数和发布解析。", "人员不再显示"],
    ["手动拆分后成员变化", "保留拆分意图；以当前有效成员和当前范围来源判断是否重新补齐。", "继续回显明确人员"],
    ["同人被其他范围覆盖", "当前范围取消不移除其他来源；人数不一定减少。", "已由其他范围选中"],
    ["明确人员同时被范围覆盖", "发布按 personId 合并；删除明确人员不影响范围覆盖。", "人数保持或按并集更新"],
    ["选择器异常关闭", "放弃临时副本，正式选择不变。", "返回原汇总页"],
    ["组织在发布前变化", "发布事务读取最新目录版本并重新解析。", "显示最终接收人数；失效则阻止"],
    ["发布后组织变化", "不修改 recipientSnapshot。", "已发布接收对象保持不变"],
    ["发布并发或重复请求", "以 idempotencyKey 保证同一发布只写入一次快照。", "返回同一发布结果"],
], [1.6, 4.25, 0.95], 7.85)


h1("7. 数据口径与埋点需求")
h2("7.1 编辑态数据")
table(["字段", "类型", "必填", "定义"], [
    ["level", "enum", "是", "scope 或 person"],
    ["scopeType", "enum", "范围是", "department 或 workgroup"],
    ["scopeId", "string", "范围是", "当前组织内稳定范围 ID"],
    ["scopeName", "string", "范围是", "选择时名称快照，仅用于回显"],
    ["selectionMode", "enum", "是", "范围为 dynamic；人员为 explicit"],
    ["personId", "string", "人员是", "人员稳定唯一 ID"],
    ["name", "string", "人员是", "人员姓名快照"],
    ["category", "enum", "人员是", "department 或 workgroup"],
    ["sourceScopeIds", "array<string>", "人员是", "产生该人员选择的范围来源；显式选择允许为空"],
    ["recipientManualPartialScopes", "array<string>", "否", "格式 scopeType:scopeId，保存手动拆分抑制状态"],
], [2.25, 1.0, 0.75, 2.8], 7.8, {2})
para('草稿示例：{ "recipientSelections": [{ "level": "scope", "scopeType": "department", "scopeId": "bureau-office", "scopeName": "办公室", "selectionMode": "dynamic" }, { "level": "person", "personId": "p-1008", "name": "李静", "category": "workgroup", "sourceScopeIds": ["project-team"] }], "recipientManualPartialScopes": ["workgroup:project-team"] }。')
h2("7.2 完整覆盖与人数算法")
bullet("有效成员集合 E(scope) = 该范围内 status=active 且 canReceive=true 的 personId 集合。E 为空时范围不可选择。")
bullet("当前范围覆盖集合 C(scope) = 当前 scope 对象解析人员 ∪ sourceScopeIds 包含当前 scopeId 的人员 ∪ 当前范围内的显式人员。")
bullet("未标记手动拆分时，E(scope) ⊆ 全部已选人员并集即可自动归一化；已标记手动拆分时，必须 E(scope) ⊆ 当前范围来源选择，才允许恢复。")
bullet("模块人数 = size(distinct personId(resolve(scopes) ∪ explicitPeople))。范围交叉、跨来源和重复选择均不得重复计数。")
h2("7.3 发布请求与响应")
para('发布请求至少包含：portalType、orgId、surveyId、directoryVersion、idempotencyKey、recipientSelections、recipientManualPartialScopes。scopeName、name 和前端 count 不作为授权或人数依据。')
para('发布响应示例：{ "surveyId": "S-20260905-01", "recipientSnapshot": [{ "personId": "p-1001", "name": "王明", "sourceScopeIds": ["bureau-office", "project-team"], "selectionMode": "snapshot", "snapshotAt": "2026-09-05T21:00:00+08:00", "frozen": true }], "recipientCount": 4, "directoryVersion": "org-20260905-2100" }。')
h2("7.4 后端实现约束")
bullet("发布服务必须根据登录态和 orgId 对 scopeId、personId 二次鉴权；越权返回 403，失效范围返回业务错误并列出 invalidScopeIds。")
bullet("范围解析、人员合并、快照写入和问卷状态变更必须在同一发布事务中完成；幂等键冲突时返回首次成功结果。")
bullet("组织目录缓存键至少包含 portalType、orgId、permissionScope、directoryVersion；权限或组织版本变化后不得复用旧结果。")
bullet("范围对象仅存在于草稿、模板和发布请求；已发布问卷的接收、填写资格与统计读取 recipientSnapshot，不再实时解析范围。")
h2("7.5 埋点")
table(["事件", "触发", "属性"], [
    ["recipient_scope_select", "选择完整部门或工作组", "portal,device,scope_type,scope_id,active_count"],
    ["recipient_scope_unselect", "取消或删除完整范围", "portal,device,scope_type,scope_id,module_selected_count"],
    ["recipient_scope_convert_to_people", "完整范围取消一人并手动拆分", "scope_type,scope_id,person_id,remaining_count"],
    ["recipient_scope_normalize", "逐人补齐并恢复范围", "scope_type,scope_id,active_count,trigger"],
    ["recipient_person_select", "选择具体人员", "portal,device,category,person_id,source_scope_id"],
    ["recipient_person_unselect", "取消具体人员", "portal,device,category,person_id,source_scope_id"],
    ["recipient_picker_confirm", "确认选择器临时副本", "portal,device,category,module_selected_count"],
    ["recipient_picker_cancel", "取消或返回", "portal,device,category,close_type"],
    ["recipient_scope_invalid", "校验到失效或越权范围", "portal,scope_type,scope_id,reason"],
    ["recipient_publish_snapshot", "发布成功固化快照", "survey_id,recipient_count,scope_count,directory_version"],
], [2.2, 2.1, 2.5], 7.8)
para("埋点不得上传姓名、手机号、搜索原文或答卷内容；personId 应按平台隐私规范脱敏。")


h1("8. 风险、依赖与限制")
table(["类型", "事项", "影响", "控制措施"], [
    ["依赖", "组织目录稳定 ID、成员状态与接收权限", "完整覆盖、人数及发布快照错误", "目录版本化；服务端二次校验"],
    ["风险", "只按全局覆盖恢复手动拆分范围", "其他范围成员误触发自动合并", "手动拆分必须按当前来源补齐"],
    ["风险", "删除标签时删除人员实体", "其他来源被误删，人数下降", "范围与人员按来源级删除"],
    ["风险", "发布后继续动态解析", "应提交人数与历史统计漂移", "发布快照 frozen=true；读取只认快照"],
    ["风险", "前后端目录版本不一致", "界面人数与发布结果不一致", "响应返回最终人数和 directoryVersion"],
    ["限制", "学校、家长和班级未改造", "通用组件误套 scope 规则", "仅 department/workgroup 进入本算法"],
], [0.65, 1.8, 2.15, 2.2], 7.8)


h1("9. 验收标准")
acceptance = [
    ["完整部门", "PC 与移动端整选部门后只回显部门名称，人数等于当前有效可接收人员去重数。", "P0"],
    ["完整工作组", "教育局 PC 与移动端整选工作组后只回显工作组名称；空工作组不可选择。", "P0"],
    ["部分人员", "任一范围未选满时仅回显具体人员姓名，不出现该范围名称。", "P0"],
    ["逐人选满", "逐人覆盖完整范围后自动归一化为范围标签，人员标签不重复展示。", "P0"],
    ["手动拆分", "完整范围取消一人仅拆分当前范围；其他完整范围保持；显示其他来源提示。", "P0"],
    ["重新补齐", "以当前范围来源补齐全部成员后清除拆分标识并恢复范围名称。", "P0"],
    ["交叉范围", "所有完整范围名称均显示；重叠人员在模块人数和发布快照中均只计一次。", "P0"],
    ["来源删除", "删除范围标签或人员标签只删除对应来源，其他覆盖关系保持。", "P0"],
    ["选择事务", "PC 弹窗与移动子页确认提交；取消、关闭、遮罩、Escape、返回均回滚。", "P0"],
    ["草稿模板", "保存并恢复 scope、person、sourceScopeIds 和手动拆分标识，跨端回显一致。", "P0"],
    ["有效成员", "停用或 canReceive=false 人员不参与完整判断、人数及快照；零成员范围无效。", "P0"],
    ["失效范围", "删除、越权或空范围在发布前被识别并阻止，不静默发布零人问卷。", "P0"],
    ["发布即时解析", "发布前新增或移除成员会反映在最终快照；同一 personId 全局去重。", "P0"],
    ["发布冻结", "发布成功后修改部门或工作组成员，已发布 recipientSnapshot 与接收人数保持不变。", "P0"],
    ["PRD一致性", "原型 prd、本文档、自动化断言和界面文案使用相同选择与发布口径。", "P0"],
]
table(["验收维度", "验收标准", "优先级"], acceptance, [1.45, 4.6, 0.75], 8.0, {2})
h2("9.1 上线前自检清单")
bullet("九个一级章节齐全且顺序固定；功能需求列表使用规定的五列表头。")
bullet("教育局 PC、学校 PC、教育局移动端、学校移动端均完成整选、部分选择、拆分与补齐检查。")
bullet("部门、工作组交叉覆盖、人数去重、其他范围标识、标签删除与选择器回滚均通过。")
bullet("草稿、模板、失效范围、发布最新成员解析和发布后快照冻结均通过自动化测试。")
bullet("JavaScript、CSS、HTML、控制台、常见移动宽度和 PC 弹窗布局均完成检查。")
bullet("DOCX 已嵌入最新交互截图，并通过全页渲染检查，无截断、溢出、空白异常或跨页表头缺失。")


doc.core_properties.title = "教育局学校问卷调研应用 PRD V1.6"
doc.core_properties.subject = "部门与工作组范围回显、手动拆分及发布人员快照"
doc.core_properties.author = "产品设计"
doc.save(OUT)
print(OUT)

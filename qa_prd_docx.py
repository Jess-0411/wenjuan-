from pathlib import Path
import os
from zipfile import ZipFile

from docx import Document


ROOT = Path(__file__).resolve().parent
DOCX = ROOT / os.environ.get("PRD_OUTPUT_NAME", "教育局学校问卷调研应用PRD_V1.6_终版.docx")
EXPECTED_HEADINGS = [
    "1. 背景与目标",
    "2. 用户与使用场景",
    "3. 需求范围",
    "4. 功能需求列表",
    "5. 核心流程与交互说明",
    "6. 异常场景与边界条件",
    "7. 数据口径与埋点需求",
    "8. 风险、依赖与限制",
    "9. 验收标准",
]
FUNCTION_HEADERS = ["功能模块", "功能点", "需求描述", "优先级（P0 / P1 / P2）", "备注说明"]


document = Document(DOCX)
headings = [p.text.strip() for p in document.paragraphs if p.style.name == "Heading 1"]
assert headings == EXPECTED_HEADINGS, f"一级章节不符合 PRD-V1：{headings}"

functional_tables = []
for target in document.tables:
    assert target.rows and target.columns, "发现空表格"
    width = len(target.columns)
    assert all(len(row.cells) == width for row in target.rows), "表格列数不一致"
    first_row = [cell.text.strip() for cell in target.rows[0].cells]
    if first_row == FUNCTION_HEADERS:
        functional_tables.append(target)
assert len(functional_tables) == 1, f"固定五列表格数量错误：{len(functional_tables)}"
assert len(functional_tables[0].rows) >= 16, "功能需求列表内容不足"

assert len(document.inline_shapes) == 5, f"交互截图数量错误：{len(document.inline_shapes)}"
all_text = "\n".join(p.text for p in document.paragraphs)
for target in document.tables:
    all_text += "\n" + "\n".join(cell.text for row in target.rows for cell in row.cells)

required_phrases = [
    "完整部门",
    "完整工作组",
    "手动拆分",
    "已由其他范围选中",
    "recipientManualPartialScopes",
    "snapshotAt",
    "frozen=true",
    "发布后组织变化",
]
for phrase in required_phrases:
    assert phrase in all_text, f"缺少关键口径：{phrase}"

forbidden_phrases = [
    "逐人选满不得自动",
    "发布时原样写入 recipientSnapshot",
    "整个收集期按最新有效成员解析",
    "工作组继续逐人选择或全选本组，不保存动态范围",
]
for phrase in forbidden_phrases:
    assert phrase not in all_text, f"存在冲突旧口径：{phrase}"

with ZipFile(DOCX) as archive:
    members = set(archive.namelist())
    media = [name for name in members if name.startswith("word/media/")]
    assert len(media) == 5, f"DOCX 媒体文件数量错误：{len(media)}"
    xml = archive.read("word/document.xml").decode("utf-8")
    assert "w:tblHeader" in xml, "跨页表格未设置重复表头"
    assert "w:cantSplit" in xml, "表格行未设置禁止跨页拆分"
    assert "<w:ins " not in xml and "<w:del " not in xml, "DOCX 存在未接受修订"

print(f"PRD DOCX structural QA passed: {len(document.paragraphs)} paragraphs, {len(document.tables)} tables, {len(media)} images")
